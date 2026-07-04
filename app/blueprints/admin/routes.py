"""Rutas del panel de administración."""
from __future__ import annotations

import io
import json

from flask import (
    Response, flash, redirect, render_template, request, send_file,
    stream_with_context, url_for, current_app,
)
from flask_login import login_required

from app.extensions import db
from app.models import (
    AdminUser, AlertaLog, Configuracion, CronLog,
    Franquiciado, Paquete,
)
from app.models.excel_import import ExcelImport

from .forms import ConfiguracionForm, FranquiciadoForm, ImportarWaybillsForm, ImportarExcelForm
from . import admin_bp


# ── Dashboard ────────────────────────────────────────────────────────────────

@admin_bp.route("/")
@login_required
def dashboard():
    total_franquiciados = Franquiciado.query.count()
    activos             = Franquiciado.query.filter_by(activo=True).count()
    total_paquetes      = Paquete.query.count()
    pendientes          = Paquete.query.filter_by(estado="pendiente").count()
    ultimo_cron         = CronLog.query.order_by(CronLog.ejecutado_at.desc()).first()
    return render_template(
        "admin/dashboard.html",
        total_franquiciados=total_franquiciados,
        activos=activos,
        total_paquetes=total_paquetes,
        pendientes=pendientes,
        ultimo_cron=ultimo_cron,
    )


# ── Franquiciados ────────────────────────────────────────────────────────────

@admin_bp.route("/franquiciados")
@login_required
def franquiciados_lista():
    franquiciados = Franquiciado.query.order_by(Franquiciado.nombre).all()
    return render_template("admin/franquiciados/lista.html", franquiciados=franquiciados)


@admin_bp.route("/franquiciados/nuevo", methods=["GET", "POST"])
@login_required
def franquiciados_nuevo():
    form = FranquiciadoForm()
    if form.validate_on_submit():
        if not form.jt_pass.data:
            flash("La contraseña JMS es obligatoria al crear un franquiciado.", "danger")
            return render_template("admin/franquiciados/form.html", form=form, modo="nuevo")

        fq = Franquiciado(
            nombre                =form.nombre.data.strip(),
            jt_user               =form.jt_user.data.strip(),
            jt_pass               =form.jt_pass.data,
            wa_grupo_id           =form.wa_grupo_id.data.strip(),
            wa_status_grupo_id    =form.wa_status_grupo_id.data.strip() or None,
            textmebot_api_key     =form.textmebot_api_key.data.strip(),
            activo                =form.activo.data,
            notas                 =form.notas.data or None,
            horas_total_entrega   =form.horas_total_entrega.data,
            horas_primera_gestion =form.horas_primera_gestion.data,
            horas_entre_gestiones =form.horas_entre_gestiones.data,
        )
        db.session.add(fq)
        db.session.commit()
        flash(f"Franquiciado «{fq.nombre}» creado correctamente.", "success")
        return redirect(url_for("admin.franquiciados_detalle", fq_id=fq.id))

    return render_template("admin/franquiciados/form.html", form=form, modo="nuevo")


@admin_bp.route("/franquiciados/<int:fq_id>")
@login_required
def franquiciados_detalle(fq_id: int):
    from datetime import datetime, timezone, timedelta
    PERU_UTC_OFFSET = timedelta(hours=5)

    fq          = db.get_or_404(Franquiciado, fq_id)
    page        = request.args.get("page", 1, type=int)
    estado      = request.args.get("estado", "")
    campo_fecha = request.args.get("campo_fecha", "")
    fecha_desde = request.args.get("fecha_desde", "")
    fecha_hasta = request.args.get("fecha_hasta", "")
    q_waybill   = request.args.get("q", "").strip()

    q_pkg = Paquete.query.filter_by(franquiciado_id=fq_id)
    if q_waybill:
        q_pkg = q_pkg.filter(Paquete.waybill_no.ilike(f"%{q_waybill}%"))
    if estado:
        q_pkg = q_pkg.filter_by(estado=estado)

    # Filtros de fecha
    _campo_map = {
        "recojo":    Paquete.fecha_recojo,
        "gestion":   Paquete.ultima_gestion_at,
        "importado": Paquete.created_at,
    }
    _col = _campo_map.get(campo_fecha)
    if _col is not None:
        if fecha_desde:
            try:
                dt_desde = datetime.strptime(fecha_desde, "%Y-%m-%d")
                q_pkg = q_pkg.filter(_col >= dt_desde)
            except ValueError:
                pass
        if fecha_hasta:
            try:
                dt_hasta = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
                q_pkg = q_pkg.filter(_col < dt_hasta)
            except ValueError:
                pass

    q_pkg      = q_pkg.order_by(Paquete.created_at.desc())
    pagination = q_pkg.paginate(page=page, per_page=50, error_out=False)
    alertas    = AlertaLog.query.filter_by(franquiciado_id=fq_id).order_by(
        AlertaLog.enviado_at.desc()
    ).limit(20).all()

    def _to_peru(dt):
        """Convierte datetime UTC a hora Perú (UTC-5). Devuelve None si dt es None."""
        if dt is None:
            return None
        return dt - PERU_UTC_OFFSET

    cfg = Configuracion.get()

    return render_template(
        "admin/franquiciados/detalle.html",
        fq=fq,
        paquetes=pagination.items,
        pagination=pagination,
        estado_filtro=estado,
        campo_fecha=campo_fecha,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        q_waybill=q_waybill,
        alertas=alertas,
        now=datetime.now(timezone.utc),
        last_wa_import_peru=_to_peru(fq.last_wa_import_at),
        last_wa_estado_peru=_to_peru(fq.last_wa_estado_at),
        last_tracking_peru=_to_peru(fq.last_tracking_at),
        cfg=cfg,
    )


@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/exportar")
@login_required
def paquetes_exportar(fq_id: int):
    """Descarga todos los paquetes filtrados como Excel."""
    import openpyxl
    from datetime import datetime, timedelta

    fq          = db.get_or_404(Franquiciado, fq_id)
    estado      = request.args.get("estado", "")
    campo_fecha = request.args.get("campo_fecha", "")
    fecha_desde = request.args.get("fecha_desde", "")
    fecha_hasta = request.args.get("fecha_hasta", "")
    q_waybill   = request.args.get("q", "").strip()

    q_pkg = Paquete.query.filter_by(franquiciado_id=fq_id)
    if q_waybill:
        q_pkg = q_pkg.filter(Paquete.waybill_no.ilike(f"%{q_waybill}%"))
    if estado:
        q_pkg = q_pkg.filter_by(estado=estado)

    _campo_map = {
        "recojo":    Paquete.fecha_recojo,
        "gestion":   Paquete.ultima_gestion_at,
        "importado": Paquete.created_at,
    }
    _col = _campo_map.get(campo_fecha)
    if _col is not None:
        if fecha_desde:
            try:
                q_pkg = q_pkg.filter(_col >= datetime.strptime(fecha_desde, "%Y-%m-%d"))
            except ValueError:
                pass
        if fecha_hasta:
            try:
                dt_hasta = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
                q_pkg = q_pkg.filter(_col < dt_hasta)
            except ValueError:
                pass

    paquetes = q_pkg.order_by(Paquete.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paquetes"

    headers = [
        "Waybill", "Estado", "Int. fallidos",
        "Fecha recojo", "Últ. intento fallido", "Últ. gestión",
        "Destinatario", "Teléfono", "Provincia", "Ciudad",
        "Dirección", "Peso (kg)", "Tipo mercancía", "Modo pago",
        "Origen pedido", "Importado",
    ]
    ws.append(headers)

    # Estilo de cabecera
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill("solid", fgColor="0B2D2A")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    def _fmt(dt):
        return dt.strftime("%d/%m/%Y %H:%M") if dt else ""

    for pkg in paquetes:
        ws.append([
            pkg.waybill_no,
            pkg.estado,
            pkg.n_intentos,
            _fmt(pkg.fecha_recojo),
            _fmt(pkg.ultimo_intento_at),
            _fmt(pkg.ultima_gestion_at),
            pkg.destinatario_nombre or "",
            pkg.destinatario_telefono or "",
            pkg.destinatario_provincia or "",
            pkg.destinatario_ciudad or "",
            pkg.destinatario_direccion or "",
            pkg.peso_cobrado,
            pkg.tipo_mercancia or "",
            pkg.modo_pago or "",
            pkg.origen_pedido or "",
            _fmt(pkg.created_at),
        ])

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"paquetes_{fq.nombre.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@admin_bp.route("/franquiciados/<int:fq_id>/editar", methods=["GET", "POST"])
@login_required
def franquiciados_editar(fq_id: int):
    fq   = db.get_or_404(Franquiciado, fq_id)
    form = FranquiciadoForm(obj=fq)

    if form.validate_on_submit():
        fq.nombre                = form.nombre.data.strip()
        fq.jt_user               = form.jt_user.data.strip()
        if form.jt_pass.data:
            fq.jt_pass = form.jt_pass.data
        fq.wa_grupo_id           = form.wa_grupo_id.data.strip()
        fq.wa_status_grupo_id    = form.wa_status_grupo_id.data.strip() or None
        fq.textmebot_api_key     = form.textmebot_api_key.data.strip()
        fq.activo                = form.activo.data
        fq.notas                 = form.notas.data or None
        fq.horas_total_entrega   = form.horas_total_entrega.data
        fq.horas_primera_gestion = form.horas_primera_gestion.data
        fq.horas_entre_gestiones = form.horas_entre_gestiones.data
        # Invalidar cache de token si cambia el usuario o clave
        fq.jt_token_cache = None
        db.session.commit()
        flash(f"Franquiciado «{fq.nombre}» actualizado.", "success")
        return redirect(url_for("admin.franquiciados_detalle", fq_id=fq.id))

    return render_template("admin/franquiciados/form.html", form=form, modo="editar", fq=fq)


@admin_bp.route("/franquiciados/<int:fq_id>/toggle", methods=["POST"])
@login_required
def franquiciados_toggle(fq_id: int):
    fq = db.get_or_404(Franquiciado, fq_id)
    fq.activo = not fq.activo
    db.session.commit()
    estado = "activado" if fq.activo else "desactivado"
    flash(f"Franquiciado «{fq.nombre}» {estado}.", "info")
    return redirect(url_for("admin.franquiciados_lista"))


# ── Paquetes ─────────────────────────────────────────────────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/importar", methods=["GET", "POST"])
@login_required
def paquetes_importar(fq_id: int):
    fq   = db.get_or_404(Franquiciado, fq_id)
    form = ImportarWaybillsForm()

    if form.validate_on_submit():
        raw_lines = form.waybills.data.splitlines()
        waybills  = [l.strip() for l in raw_lines if l.strip()]

        agregados = 0
        duplicados = 0
        for wb in waybills:
            existente = Paquete.query.filter_by(
                franquiciado_id=fq_id, waybill_no=wb
            ).first()
            if existente:
                duplicados += 1
                continue
            db.session.add(Paquete(
                franquiciado_id=fq_id,
                waybill_no=wb,
                estado="pendiente",
            ))
            agregados += 1

        db.session.commit()
        flash(
            f"{agregados} paquete(s) importado(s). {duplicados} duplicado(s) ignorado(s).",
            "success",
        )
        return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))

    return render_template("admin/paquetes/importar.html", form=form, fq=fq)


# ── Importar Excel (J&T "Monitoreo entrada al puerto") ───────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/importar-excel", methods=["GET", "POST"])
@login_required
def paquetes_importar_excel(fq_id: int):
    """Sube un Excel exportado desde JMS y agrega sus waybills al franquiciado."""
    from app.services.importar_service import parse_excel_waybills, import_waybills
    from app.services.google_drive_service import google_drive_service as drive_service

    fq   = db.get_or_404(Franquiciado, fq_id)
    form = ImportarExcelForm()

    if form.validate_on_submit():
        uploaded = form.excel_file.data
        file_bytes = uploaded.read()
        filename   = uploaded.filename

        # 1. Parsear Excel
        try:
            waybills = parse_excel_waybills(file_bytes)
        except Exception as exc:
            flash(f"Error al leer el Excel: {exc}", "danger")
            return render_template("admin/paquetes/importar_excel.html", form=form, fq=fq)

        if not waybills:
            flash("El Excel no contiene waybills válidos.", "warning")
            return render_template("admin/paquetes/importar_excel.html", form=form, fq=fq)

        # 2. Subir a Google Drive (si hay credenciales)
        drive_file_id = None
        folder_id = current_app.config.get("GOOGLE_DRIVE_FOLDER_ID")
        if folder_id:
            try:
                drive_file_id = drive_service.upload_file(
                    file_bytes=file_bytes,
                    filename=filename,
                    mime_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        if filename.lower().endswith(".xlsx")
                        else "application/vnd.ms-excel"
                    ),
                    folder_id=folder_id,
                )
                flash(
                    f'Excel guardado en Drive. '
                    f'<a href="https://drive.google.com/file/d/{drive_file_id}/view" '
                    f'target="_blank">Ver archivo</a>',
                    "info",
                )
            except Exception as exc:
                current_app.logger.warning(f"Drive upload falló: {exc}")
                flash("No se pudo subir el Excel a Google Drive (credenciales no configuradas o error de red).", "warning")

        # 3. Importar a BD
        result = import_waybills(
            franquiciado_id=fq_id,
            waybills=waybills,
            filename=filename,
            drive_file_id=drive_file_id,
            import_mode="manual",
        )

        flash(
            f"{result['nuevos']} paquete(s) agregado(s). "
            f"{result['duplicados']} duplicado(s) ignorado(s). "
            f"Total en Excel: {result['total']}.",
            "success",
        )
        return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))

    return render_template("admin/paquetes/importar_excel.html", form=form, fq=fq)


# ── Sincronizar automáticamente desde el portal J&T ─────────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/sincronizar", methods=["POST"])
@login_required
def paquetes_sincronizar(fq_id: int):
    """Usa el scraper OutletMonitor para obtener waybills directamente del portal."""
    from datetime import date, timedelta
    from jt_scraper.outlet_monitor import OutletMonitor
    from jt_scraper.instance_config import JTInstanceConfig
    from app.services.importar_service import import_waybills
    from app.models import Configuracion

    fq  = db.get_or_404(Franquiciado, fq_id)
    cfg = Configuracion.get()

    instance_cfg = JTInstanceConfig(
        jt_user=fq.jt_user,
        jt_pass=fq.jt_pass,
        token_getter=fq.get_token_cache,
        token_setter=lambda v: (fq.set_token_cache(v), db.session.commit()),
    )

    try:
        monitor    = OutletMonitor(instance_cfg)
        start_date = (date.today() - timedelta(days=cfg.sync_dias_atras)).isoformat()
        end_date   = date.today().isoformat()
        waybills   = monitor.fetch_waybills(start_date, end_date, time_type=cfg.sync_time_type)

        if not waybills:
            flash("El scraper no encontró waybills en el portal.", "warning")
            return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))

        result = import_waybills(
            franquiciado_id=fq_id,
            waybills=waybills,
            filename=f"auto_sync_{end_date}.xlsx",
            import_mode="auto",
        )
        flash(
            f"Sincronización completa: {result['nuevos']} nuevo(s), "
            f"{result['duplicados']} duplicado(s). Total encontrado: {result['total']}.",
            "success",
        )
    except Exception as exc:
        current_app.logger.exception(f"Error en sincronización franq={fq_id}")
        flash(f"Error en sincronización: {exc}", "danger")

    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


# ── Comandos WhatsApp desde el panel admin ───────────────────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/wa/estado", methods=["POST"])
@login_required
def wa_estado(fq_id: int):
    """Envía el resumen de estados de paquetes al grupo WhatsApp del franquiciado."""
    from datetime import datetime, timedelta
    from app.blueprints.whatsapp.webhook import _build_estado_message
    from app.services.whatsapp_service import enviar_whatsapp

    PERU_UTC_OFFSET = timedelta(hours=5)
    fq = db.get_or_404(Franquiciado, fq_id)

    now_utc  = datetime.utcnow()
    now_peru = now_utc if current_app.debug else (now_utc - PERU_UTC_OFFSET)
    msg = _build_estado_message(fq, now_peru)

    ok, resp = enviar_whatsapp(fq.textmebot_api_key, fq.wa_grupo_id, msg)

    fq.last_wa_estado_at = now_utc
    db.session.commit()

    if ok:
        flash("Resumen de estado enviado al grupo WhatsApp correctamente.", "success")
    else:
        flash(f"Error al enviar a WhatsApp: {resp[:200]}", "danger")

    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


@admin_bp.route("/franquiciados/<int:fq_id>/wa/importar", methods=["POST"])
@login_required
def wa_importar(fq_id: int):
    """Importa paquetes desde JMS y envía el resultado al grupo WhatsApp.
    Sin rate limit — acceso exclusivo para el admin."""
    from datetime import date, datetime, timedelta
    from jt_scraper.outlet_monitor import OutletMonitor
    from jt_scraper.instance_config import JTInstanceConfig
    from app.services.importar_service import import_waybills
    from app.services.whatsapp_service import enviar_whatsapp
    from app.models import Configuracion

    fq  = db.get_or_404(Franquiciado, fq_id)
    cfg = Configuracion.get()

    instance_cfg = JTInstanceConfig(
        jt_user=fq.jt_user,
        jt_pass=fq.jt_pass,
        token_getter=fq.get_token_cache,
        token_setter=lambda v: (fq.set_token_cache(v), db.session.commit()),
    )

    try:
        monitor    = OutletMonitor(instance_cfg)
        start_date = (date.today() - timedelta(days=cfg.sync_dias_atras)).isoformat()
        end_date   = date.today().isoformat()
        waybills   = monitor.fetch_waybills(start_date, end_date, time_type=cfg.sync_time_type)

        if not waybills:
            msg = "⚠️ No se encontraron paquetes en JMS para el rango de fechas consultado."
            enviar_whatsapp(fq.textmebot_api_key, fq.wa_grupo_id, msg)
            flash("JMS no devolvió paquetes. Se notificó al grupo.", "warning")
        else:
            result = import_waybills(
                franquiciado_id=fq_id,
                waybills=waybills,
                filename=f"admin_wa_importar_{end_date}",
                import_mode="wa_admin",
            )
            wa_msg = (
                f"✅ *Importación completada (admin)*\n"
                f"• Paquetes nuevos: *{result['nuevos']}*\n"
                f"• Duplicados (ya existían): *{result['duplicados']}*\n"
                f"• Total procesados: *{result['total']}*\n\n"
                f"_Usa /estado para ver el resumen actualizado._"
            )
            enviar_whatsapp(fq.textmebot_api_key, fq.wa_grupo_id, wa_msg)
            flash(
                f"Importación completada: {result['nuevos']} nuevo(s), "
                f"{result['duplicados']} duplicado(s). Resultado enviado al grupo WA.",
                "success",
            )

        fq.last_wa_import_at = datetime.utcnow()
        db.session.commit()

    except Exception as exc:
        current_app.logger.exception(f"[admin] Error en wa_importar franq={fq_id}")
        flash(f"Error en importación: {exc}", "danger")

    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


@admin_bp.route("/paquetes/<int:pkg_id>/cancelar", methods=["POST"])
@login_required
def paquetes_cancelar(pkg_id: int):
    pkg = db.get_or_404(Paquete, pkg_id)
    fq_id = pkg.franquiciado_id
    pkg.estado = Paquete.ESTADO_CANCELADO
    db.session.commit()
    flash(f"Paquete {pkg.waybill_no} cancelado.", "info")
    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


@admin_bp.route("/paquetes/<int:pkg_id>/siniestrar", methods=["POST"])
@login_required
def paquetes_siniestrar(pkg_id: int):
    pkg = db.get_or_404(Paquete, pkg_id)
    fq_id = pkg.franquiciado_id
    pkg.estado = Paquete.ESTADO_SINIESTRADO
    db.session.commit()
    flash(f"Paquete {pkg.waybill_no} marcado como siniestrado.", "warning")
    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


# ── Actualizar tracking de un franquiciado ───────────────────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/actualizar-tracking", methods=["POST"])
@login_required
def paquetes_actualizar_tracking(fq_id: int):
    """Consulta el POD de J&T para cada paquete pendiente del franquiciado
    y actualiza fecha_recojo, n_intentos, ultimo_intento_at, ultima_gestion_at.
    Equivale a ejecutar el cron solo para este franquiciado."""
    from app.services.tracking_service import refrescar_franquiciado

    fq = db.get_or_404(Franquiciado, fq_id)

    try:
        stats = refrescar_franquiciado(fq, flask_debug=current_app.debug)
        from datetime import datetime
        fq.last_tracking_at = datetime.utcnow()
        db.session.commit()
        flash(
            f"Tracking actualizado: {stats['consultados']} paquete(s) procesado(s). "
            f"Entregados: {stats['entregados']} · Devueltos: {stats['devueltos']} · "
            f"Errores: {stats['errores']}.",
            "success" if stats["errores"] == 0 else "warning",
        )
    except Exception as exc:
        current_app.logger.exception(f"Error actualizando tracking franq={fq_id}")
        flash(f"Error al actualizar tracking: {exc}", "danger")

    return redirect(url_for("admin.franquiciados_detalle", fq_id=fq_id))


# ── Streaming SSE de progreso de tracking (lotes) ───────────────────────────

# Paquetes por lote SSE. Cada lote dura ~30–40 s, muy por debajo del
# timeout de gunicorn (120 s). El cliente JS reconecta hasta que el
# servidor responde {type:'all_done'}.
_TRACKING_BATCH = 20


@admin_bp.route("/franquiciados/<int:fq_id>/paquetes/tracking-stream")
@login_required
def paquetes_tracking_stream(fq_id: int):
    """SSE con lotes: procesa _TRACKING_BATCH paquetes por conexión.

    Query params:
        after_id (int, default 0): cursor — procesa paquetes con id > after_id
    """
    from app.services.tracking_service import refrescar_lote_stream

    fq       = db.get_or_404(Franquiciado, fq_id)
    after_id = request.args.get("after_id", 0, type=int)
    debug    = current_app.debug

    total_pending = (
        Paquete.query
        .filter_by(franquiciado_id=fq_id, estado=Paquete.ESTADO_PENDIENTE)
        .count()
    )

    batch = (
        Paquete.query
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.estado          == Paquete.ESTADO_PENDIENTE,
            Paquete.id              >  after_id,
        )
        .order_by(Paquete.id)
        .limit(_TRACKING_BATCH)
        .all()
    )

    def _generate():
        if not batch:
            yield f"data: {json.dumps({'type': 'all_done'})}\n\n"
            return

        yield f"data: {json.dumps({'type': 'start', 'total': total_pending, 'batch': len(batch)})}\n\n"

        for event in refrescar_lote_stream(fq, batch, flask_debug=debug):
            yield f"data: {json.dumps(event)}\n\n"

    return Response(
        stream_with_context(_generate()),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


# ── Cron Logs ────────────────────────────────────────────────────────────────

@admin_bp.route("/cron-logs")
@login_required
def cron_logs():
    logs = CronLog.query.order_by(CronLog.ejecutado_at.desc()).limit(100).all()
    return render_template("admin/cron_logs/lista.html", logs=logs)


# ── Alertas Log ──────────────────────────────────────────────────────────────

@admin_bp.route("/alertas")
@login_required
def alertas_lista():
    page    = request.args.get("page", 1, type=int)
    fq_id   = request.args.get("franquiciado_id", type=int)

    query = AlertaLog.query.order_by(AlertaLog.enviado_at.desc())
    if fq_id:
        query = query.filter_by(franquiciado_id=fq_id)

    pagination  = query.paginate(page=page, per_page=50, error_out=False)
    franquiciados = Franquiciado.query.order_by(Franquiciado.nombre).all()
    return render_template(
        "admin/alertas/lista.html",
        pagination=pagination,
        alertas=pagination.items,
        franquiciados=franquiciados,
        fq_id_filtro=fq_id,
    )


# ── Dashboard por franquiciado ───────────────────────────────────────────────

@admin_bp.route("/franquiciados/<int:fq_id>/dashboard")
@login_required
def franquiciados_dashboard(fq_id: int):
    from datetime import datetime, timedelta
    from sqlalchemy import func, cast
    from sqlalchemy.types import Date as SADate

    fq = db.get_or_404(Franquiciado, fq_id)

    hoy = datetime.utcnow().date()

    # Primera fecha con datos del franquiciado (recojo o entrega)
    _min_recojo = db.session.query(func.min(Paquete.fecha_recojo)).filter(
        Paquete.franquiciado_id == fq_id, Paquete.fecha_recojo.isnot(None)
    ).scalar()
    _min_entrega = db.session.query(func.min(Paquete.ultima_gestion_at)).filter(
        Paquete.franquiciado_id == fq_id, Paquete.ultima_gestion_at.isnot(None)
    ).scalar()

    def _to_date(v):
        if v is None:
            return None
        return v.date() if hasattr(v, "date") and callable(v.date) else v

    _candidates = [_to_date(_min_recojo), _to_date(_min_entrega)]
    primer_dia   = min(d for d in _candidates if d is not None) if any(_candidates) else hoy - timedelta(days=29)
    fecha_inicio = datetime(primer_dia.year, primer_dia.month, primer_dia.day)

    # ── 1. Totales por estado ────────────────────────────────────────────────
    por_estado_raw = (
        db.session.query(Paquete.estado, func.count().label("n"))
        .filter(Paquete.franquiciado_id == fq_id)
        .group_by(Paquete.estado)
        .all()
    )
    por_estado     = {r.estado: r.n for r in por_estado_raw}
    total_paquetes = sum(por_estado.values())
    tasa_entrega   = (
        round(por_estado.get("entregado", 0) / total_paquetes * 100, 1)
        if total_paquetes else 0
    )

    # ── 2. Series diarias desde primer dato ────────────────────────────────
    n_dias      = (hoy - primer_dia).days + 1
    dias_date   = [primer_dia + timedelta(days=i) for i in range(n_dias)]
    dias_iso    = [d.isoformat() for d in dias_date]
    dias_labels = [d.strftime("%d/%m") for d in dias_date]

    def _key(val):
        """Normaliza date/string → 'YYYY-MM-DD'."""
        return val.isoformat() if hasattr(val, "isoformat") else str(val)[:10]

    dia_recojo  = cast(Paquete.fecha_recojo, SADate)
    recogidos_raw = (
        db.session.query(dia_recojo.label("dia"), func.count().label("n"))
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.fecha_recojo.isnot(None),
            Paquete.fecha_recojo >= fecha_inicio,
        )
        .group_by(dia_recojo)
        .order_by(dia_recojo)
        .all()
    )
    recogidos_map  = {_key(r.dia): r.n for r in recogidos_raw}
    recogidos_serie = [recogidos_map.get(d, 0) for d in dias_iso]

    dia_entrega = cast(Paquete.ultima_gestion_at, SADate)
    entregados_raw = (
        db.session.query(dia_entrega.label("dia"), func.count().label("n"))
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.estado == "entregado",
            Paquete.ultima_gestion_at.isnot(None),
            Paquete.ultima_gestion_at >= fecha_inicio,
        )
        .group_by(dia_entrega)
        .order_by(dia_entrega)
        .all()
    )
    entregados_map  = {_key(r.dia): r.n for r in entregados_raw}
    entregados_serie = [entregados_map.get(d, 0) for d in dias_iso]

    # ── 3. Entregas por hora del día ─────────────────────────────────────────
    # ultima_gestion_at ya se almacena en hora peruana (la API JyT la devuelve
    # así, sin conversión adicional). No se aplica offset.
    ts_rows = (
        db.session.query(Paquete.ultima_gestion_at)
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.estado == "entregado",
            Paquete.ultima_gestion_at.isnot(None),
        )
        .all()
    )
    por_hora = [0] * 24
    for (dt,) in ts_rows:
        por_hora[dt.hour] += 1

    manana = sum(por_hora[6:12])
    tarde  = sum(por_hora[12:18])
    noche  = sum(por_hora[18:24]) + sum(por_hora[0:6])

    # ── 4. Top provincias ────────────────────────────────────────────────────
    provincias_raw = (
        db.session.query(Paquete.destinatario_provincia, func.count().label("n"))
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.destinatario_provincia.isnot(None),
        )
        .group_by(Paquete.destinatario_provincia)
        .order_by(func.count().desc())
        .limit(8)
        .all()
    )
    provincias_labels = [r.destinatario_provincia for r in provincias_raw]
    provincias_values = [r.n for r in provincias_raw]

    # ── 5. Por origen ────────────────────────────────────────────────────────
    origenes_raw = (
        db.session.query(Paquete.origen_pedido, func.count().label("n"))
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.origen_pedido.isnot(None),
        )
        .group_by(Paquete.origen_pedido)
        .order_by(func.count().desc())
        .limit(6)
        .all()
    )
    origenes_labels = [r.origen_pedido for r in origenes_raw]
    origenes_values = [r.n for r in origenes_raw]

    # ── 6. Intentos fallidos (paquetes entregados) ───────────────────────────
    intentos_raw = (
        db.session.query(Paquete.n_intentos, func.count().label("n"))
        .filter(
            Paquete.franquiciado_id == fq_id,
            Paquete.estado == "entregado",
        )
        .group_by(Paquete.n_intentos)
        .order_by(Paquete.n_intentos)
        .all()
    )
    intentos_labels = [str(r.n_intentos) for r in intentos_raw]
    intentos_values = [r.n for r in intentos_raw]

    return render_template(
        "admin/franquiciados/dashboard.html",
        fq=fq,
        total_paquetes=total_paquetes,
        por_estado=por_estado,
        tasa_entrega=tasa_entrega,
        dias_labels=dias_labels,
        evolucion_desde=primer_dia.strftime("%d/%m/%Y"),
        recogidos_serie=recogidos_serie,
        entregados_serie=entregados_serie,
        por_hora=por_hora,
        manana=manana,
        tarde=tarde,
        noche=noche,
        provincias_labels=provincias_labels,
        provincias_values=provincias_values,
        hay_provincias=bool(provincias_labels),
        origenes_labels=origenes_labels,
        origenes_values=origenes_values,
        hay_origenes=bool(origenes_labels),
        intentos_labels=intentos_labels,
        intentos_values=intentos_values,
        hay_intentos=bool(intentos_labels),
        hay_entregas=bool(ts_rows),
    )


# ── Configuración ────────────────────────────────────────────────────────────

@admin_bp.route("/configuracion", methods=["GET", "POST"])
@login_required
def configuracion():
    cfg  = Configuracion.get()
    form = ConfiguracionForm(obj=cfg)

    if form.validate_on_submit():
        cfg.umbral_dia            = form.umbral_dia.data
        cfg.umbral_22             = form.umbral_22.data
        cfg.umbral_23             = form.umbral_23.data
        cfg.hora_inicio           = form.hora_inicio.data
        cfg.hora_fin              = form.hora_fin.data
        cfg.delay_whatsapp        = form.delay_whatsapp.data
        cfg.sync_dias_atras       = form.sync_dias_atras.data
        cfg.sync_time_type        = form.sync_time_type.data
        db.session.commit()
        flash("Configuración guardada.", "success")
        return redirect(url_for("admin.configuracion"))

    return render_template("admin/configuracion/form.html", form=form, cfg=cfg)
