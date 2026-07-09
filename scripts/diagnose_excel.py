"""
Diagnóstico del parser de Excel.
Uso: python scripts/diagnose_excel.py "ruta\al\archivo.xlsx"
"""
import sys
import io
import unicodedata
from openpyxl import load_workbook


_GUIA_ALIASES = {
    "numerodeguia", "numerodeguias", "waybillno", "waybill",
    "trackingno", "codigodeguia", "guia", "tracking",
    "numeroguia", "numguia", "numguias",
}


def _normalize_old(s: str) -> str:
    """Con el bug actual (elimina 'n')."""
    s = unicodedata.normalize("NFD", str(s).lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    for ch in (" ", ".", "_", "-", "#", "°", "º", "n"):
        s = s.replace(ch, "")
    return s


def _normalize_fixed(s: str) -> str:
    """Sin el bug (no elimina 'n')."""
    s = unicodedata.normalize("NFD", str(s).lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    for ch in (" ", ".", "_", "-", "#", "°", "º"):
        s = s.replace(ch, "")
    return s


def main(filepath: str):
    with open(filepath, "rb") as f:
        file_bytes = f.read()

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    print(f"Hoja activa : {ws.title}")

    print("=" * 70)
    print("PRIMERAS 5 FILAS (primeras 15 columnas):")
    print("=" * 70)
    for i, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
        vals = list(row[:15])
        print(f"  Fila {i+1:2d}: {vals}")

    print("\n" + "=" * 70)
    print("ANÁLISIS DE ENCABEZADOS (primeras 10 filas):")
    print("=" * 70)
    wb2 = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws2 = wb2.active

    header_row_idx = None
    guia_col_idx = None

    for row_idx, row in enumerate(ws2.iter_rows(max_row=10, values_only=True)):
        for col_idx, cell_val in enumerate(row):
            if cell_val is None:
                continue
            old = _normalize_old(str(cell_val))
            new = _normalize_fixed(str(cell_val))
            match_old = old in _GUIA_ALIASES or any(a in old for a in _GUIA_ALIASES)
            match_new = new in _GUIA_ALIASES or any(a in new for a in _GUIA_ALIASES)
            if match_old or match_new:
                print(f"  Fila {row_idx+1}, Col {col_idx} — valor: '{cell_val}'")
                print(f"    normalize_old  → '{old}'  match={match_old}")
                print(f"    normalize_fixed→ '{new}'  match={match_new}")
                if header_row_idx is None and match_old:
                    header_row_idx = row_idx
                    guia_col_idx = col_idx
                    print(f"    *** COLUMNA DETECTADA por código actual ***")

    wb2.close()

    if guia_col_idx is None:
        print("\n⚠ No se detectó columna de guía con el código actual.")
        wb.close()
        return

    print(f"\n{'='*70}")
    print(f"PRIMERAS 10 CELDAS DE DATOS (col={guia_col_idx}, min_row={header_row_idx+2}):")
    print(f"{'='*70}")
    wb3 = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws3 = wb3.active
    count = 0
    for row in ws3.iter_rows(min_row=header_row_idx + 2, values_only=True):
        val = row[guia_col_idx] if guia_col_idx < len(row) else None
        print(f"  val={repr(val)}  type={type(val).__name__}")
        count += 1
        if count >= 10:
            break
    wb3.close()
    wb.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python scripts/diagnose_excel.py \"ruta\\al\\archivo.xlsx\"")
        sys.exit(1)
    main(sys.argv[1])
